import re

with open('routes/dashboard.py', 'r', encoding='utf-8') as f:
    content = f.read()

# 1. Update index() variables
old_vars = '''    periodo = request.args.get('periodo', 'mes')
    hoy = get_local_now()
    inicio_semana = (hoy - timedelta(days=hoy.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    inicio_mes = hoy.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    
    if periodo == 'dia':
        inicio_periodo = hoy.replace(hour=0, minute=0, second=0, microsecond=0)
    elif periodo == 'semana':
        inicio_periodo = inicio_semana
    else:  # mes
        inicio_periodo = inicio_mes'''

new_vars = '''    periodo = request.args.get('periodo', 'mes')
    mes_historico = request.args.get('mes_historico')
    hoy = get_local_now()
    inicio_semana = (hoy - timedelta(days=hoy.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    inicio_mes = hoy.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    
    fin_periodo = None
    
    if periodo == 'historico' and mes_historico:
        year, month = map(int, mes_historico.split('-'))
        inicio_periodo = hoy.replace(year=year, month=month, day=1, hour=0, minute=0, second=0, microsecond=0)
        if month == 12:
            fin_periodo = hoy.replace(year=year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
        else:
            fin_periodo = hoy.replace(year=year, month=month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
    elif periodo == 'dia':
        inicio_periodo = hoy.replace(hour=0, minute=0, second=0, microsecond=0)
    elif periodo == 'semana':
        inicio_periodo = inicio_semana
    else:  # mes
        inicio_periodo = inicio_mes

    filtros_fecha = [Factura.fecha_emision >= inicio_periodo]
    if fin_periodo:
        filtros_fecha.append(Factura.fecha_emision < fin_periodo)
'''
content = content.replace(old_vars, new_vars)

# 2. Update queries
old_q1 = '''    top_clientes = db.session.query(
        Cliente.nombre_completo,
        func.sum(Paquete.peso).label('total_libras'),
        func.count(Paquete.id).label('total_paquetes')
    ).join(Factura, Cliente.id == Factura.cliente_id)\\
     .join(Paquete, Factura.id == Paquete.factura_id)\\
     .filter(Cliente.activo == True, Factura.estado.in_(['finalizada', 'pagada']), Factura.fecha_emision >= inicio_periodo)\\
     .group_by(Cliente.id)\\
     .order_by(func.sum(Paquete.peso).desc())\\
     .limit(5).all()'''

new_q1 = '''    top_clientes = db.session.query(
        Cliente.nombre_completo,
        func.sum(Paquete.peso).label('total_libras'),
        func.count(Paquete.id).label('total_paquetes')
    ).join(Factura, Cliente.id == Factura.cliente_id)\\
     .join(Paquete, Factura.id == Paquete.factura_id)\\
     .filter(Cliente.activo == True, Factura.estado.in_(['finalizada', 'pagada']), *filtros_fecha)\\
     .group_by(Cliente.id)\\
     .order_by(func.sum(Paquete.peso).desc())\\
     .limit(5).all()'''
content = content.replace(old_q1, new_q1)

old_q2 = '''    aereos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'aereo',
        Factura.estado.in_(['finalizada', 'pagada']),
        Factura.fecha_emision >= inicio_periodo
    ).count()

    maritimos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'maritimo',
        Factura.estado.in_(['finalizada', 'pagada']),
        Factura.fecha_emision >= inicio_periodo
    ).count()'''

new_q2 = '''    aereos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'aereo',
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).count()

    maritimos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'maritimo',
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).count()'''
content = content.replace(old_q2, new_q2)

old_loop = '''        # Últimos 6 meses
        for i in range(5, -1, -1):
            target_month = hoy.month - i
            target_year = hoy.year
            while target_month <= 0:
                target_month += 12
                target_year -= 1
                
            inicio = hoy.replace(year=target_year, month=target_month, day=1, hour=0, minute=0, second=0, microsecond=0)
            if target_month == 12:
                fin = hoy.replace(year=target_year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
            else:
                fin = hoy.replace(year=target_year, month=target_month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
                
            total = calcular_ganancia_facturas(inicio, fin)
            tendencia_labels.append(inicio.strftime('%b %Y'))
            tendencia_valores.append(float(total))'''

new_loop = '''        # Últimos 6 meses hasta inicio_periodo
        base_date = inicio_periodo if periodo == 'historico' else hoy
        for i in range(5, -1, -1):
            target_month = base_date.month - i
            target_year = base_date.year
            while target_month <= 0:
                target_month += 12
                target_year -= 1
                
            inicio = base_date.replace(year=target_year, month=target_month, day=1, hour=0, minute=0, second=0, microsecond=0)
            if target_month == 12:
                fin = base_date.replace(year=target_year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
            else:
                fin = base_date.replace(year=target_year, month=target_month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
                
            total = calcular_ganancia_facturas(inicio, fin)
            tendencia_labels.append(inicio.strftime('%b %Y'))
            tendencia_valores.append(float(total))'''
content = content.replace(old_loop, new_loop)

old_return = '''    return render_template('dashboard/index.html',
        periodo=periodo,
        ganancias_semana=ganancias_semana,'''
new_return = '''    return render_template('dashboard/index.html',
        periodo=periodo,
        mes_historico=mes_historico,
        ganancias_semana=ganancias_semana,'''
content = content.replace(old_return, new_return)

excel_endpoint = '''
@dashboard_bp.route('/exportar-excel')
@login_required
def exportar_excel():
    if current_user.rol != 'admin':
        flash('No tienes permisos.', 'error')
        return redirect(url_for('paquetes.index'))
        
    from flask import make_response
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    periodo = request.args.get('periodo', 'mes')
    mes_historico = request.args.get('mes_historico')
    hoy = get_local_now()
    inicio_semana = (hoy - timedelta(days=hoy.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    inicio_mes = hoy.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    
    fin_periodo = None
    
    if periodo == 'historico' and mes_historico:
        year, month = map(int, mes_historico.split('-'))
        inicio_periodo = hoy.replace(year=year, month=month, day=1, hour=0, minute=0, second=0, microsecond=0)
        if month == 12:
            fin_periodo = hoy.replace(year=year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
        else:
            fin_periodo = hoy.replace(year=year, month=month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
        titulo_periodo = f'Histórico ({mes_historico})'
    elif periodo == 'dia':
        inicio_periodo = hoy.replace(hour=0, minute=0, second=0, microsecond=0)
        titulo_periodo = 'Diario'
    elif periodo == 'semana':
        inicio_periodo = inicio_semana
        titulo_periodo = 'Semanal'
    else:  # mes
        inicio_periodo = inicio_mes
        titulo_periodo = 'Mensual'
        
    filtros_fecha = [Factura.fecha_emision >= inicio_periodo]
    if fin_periodo:
        filtros_fecha.append(Factura.fecha_emision < fin_periodo)

    # 1. KPIs Globales
    ganancia_periodo = calcular_ganancia_facturas(inicio_periodo, fin_periodo)
    
    # We only count those that were invoiced in this period
    aereos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'aereo',
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).count()
    maritimos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'maritimo',
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).count()

    total_clientes = Cliente.query.filter_by(activo=True).count()
    total_paquetes = Paquete.query.count()

    wb = openpyxl.Workbook()
    
    # --- HOJA 1: RESUMEN ---
    ws1 = wb.active
    ws1.title = "Resumen"
    header_fill = PatternFill(start_color='3D5BA0', end_color='3D5BA0', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')

    ws1.append([f"RESUMEN {titulo_periodo.upper()}"])
    ws1['A1'].font = Font(bold=True, size=14, color='3D5BA0')
    ws1.append([])
    
    resumen_data = [
        ["Indicador", "Valor"],
        ["Ganancia del Periodo ($)", round(ganancia_periodo, 2)],
        ["Paquetes Aéreos (Facturados)", aereos],
        ["Paquetes Marítimos (Facturados)", maritimos],
        ["Total Clientes Activos (Global)", total_clientes],
        ["Total Paquetes en Sistema", total_paquetes]
    ]
    
    for row_idx, row_data in enumerate(resumen_data, 3):
        ws1.append(row_data)
        if row_idx == 3:
            ws1.cell(row=row_idx, column=1).fill = header_fill
            ws1.cell(row=row_idx, column=1).font = header_font
            ws1.cell(row=row_idx, column=2).fill = header_fill
            ws1.cell(row=row_idx, column=2).font = header_font

    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['B'].width = 15

    # --- HOJA 2: GRAFICA ---
    ws2 = wb.create_sheet(title="Evolución (Gráfica)")
    ws2.append(["Periodo", "Ganancia Neta ($)"])
    ws2['A1'].fill = header_fill; ws2['A1'].font = header_font
    ws2['B1'].fill = header_fill; ws2['B1'].font = header_font

    if periodo == 'dia':
        for i in range(6, -1, -1):
            dia = hoy - timedelta(days=i)
            inicio = dia.replace(hour=0, minute=0, second=0, microsecond=0)
            fin = inicio + timedelta(days=1)
            total = calcular_ganancia_facturas(inicio, fin)
            lbl = inicio.strftime('%d %b')
            ws2.append([lbl if i != 0 else "Hoy", total])
    elif periodo == 'semana':
        for i in range(3, -1, -1):
            inicio = inicio_semana - timedelta(days=7 * i)
            fin = inicio + timedelta(days=7)
            total = calcular_ganancia_facturas(inicio, fin)
            lbl = f"{inicio.strftime('%d %b')} - {(fin - timedelta(days=1)).strftime('%d %b')}"
            ws2.append([lbl if i != 0 else "Esta Semana", total])
    else:
        base_date = inicio_periodo if periodo == 'historico' else hoy
        for i in range(5, -1, -1):
            target_month = base_date.month - i
            target_year = base_date.year
            while target_month <= 0:
                target_month += 12; target_year -= 1
            inicio = base_date.replace(year=target_year, month=target_month, day=1, hour=0, minute=0, second=0, microsecond=0)
            if target_month == 12: fin = base_date.replace(year=target_year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
            else: fin = base_date.replace(year=target_year, month=target_month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
            total = calcular_ganancia_facturas(inicio, fin)
            ws2.append([inicio.strftime('%b %Y'), total])

    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 20

    # --- HOJA 3: TOP CLIENTES ---
    ws3 = wb.create_sheet(title="Top Clientes")
    ws3.append(["Posición", "Cliente", "Total Libras", "Total Paquetes"])
    for col in ['A', 'B', 'C', 'D']:
        ws3[f'{col}1'].fill = header_fill; ws3[f'{col}1'].font = header_font

    top_clientes_q = db.session.query(
        Cliente.nombre_completo,
        func.sum(Paquete.peso).label('total_libras'),
        func.count(Paquete.id).label('total_paquetes')
    ).join(Factura, Cliente.id == Factura.cliente_id).join(Paquete, Factura.id == Paquete.factura_id).filter(
        Cliente.activo == True, Factura.estado.in_(['finalizada', 'pagada']), *filtros_fecha
    ).group_by(Cliente.id).order_by(func.sum(Paquete.peso).desc()).limit(10).all()

    for idx, (nombre, libras, paquetes) in enumerate(top_clientes_q, 1):
        ws3.append([idx, nombre, float(libras or 0), paquetes])

    ws3.column_dimensions['B'].width = 30
    ws3.column_dimensions['C'].width = 15
    ws3.column_dimensions['D'].width = 15

    # --- HOJA 4: DETALLE FACTURAS ---
    ws4 = wb.create_sheet(title="Detalle Facturas")
    headers4 = ["Factura ID", "Número", "Fecha Emisión", "Cliente", "Estado", "Total ($)", "Costo Agencia ($)", "Ganancia Neta ($)"]
    ws4.append(headers4)
    for i, _ in enumerate(headers4, 1):
        ws4.cell(row=1, column=i).fill = header_fill
        ws4.cell(row=1, column=i).font = header_font

    facturas = Factura.query.options(joinedload(Factura.paquetes)).filter(
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).order_by(Factura.fecha_emision.desc()).all()

    from models import COSTOS_AGENCIA
    for f in facturas:
        costo_ag = 0
        for p in f.paquetes:
            if getattr(p, 'categoria', 'general') in ['celular', 'laptop']:
                origen_key = p.origen if getattr(p, 'origen', None) else 'miami'
                costo_unidad = COSTOS_AGENCIA.get(origen_key, COSTOS_AGENCIA['miami']).get(p.categoria, 0)
                costo_ag += (p.cantidad or 1) * costo_unidad
            else:
                origen = p.origen if getattr(p, 'origen', None) else 'miami'
                tarifas_origen = COSTOS_AGENCIA.get(origen, COSTOS_AGENCIA['miami'])
                tipo = 'aereo' if str(p.tipo_envio).strip().lower() in ['aereo', 'aéreo'] else 'maritimo'
                costo_ag += (p.peso * tarifas_origen.get(tipo, tarifas_origen['aereo']))
        ganancia = f.total - costo_ag
        ws4.append([f.id, f.numero, f.fecha_emision.strftime('%Y-%m-%d %H:%M') if f.fecha_emision else '',
                    f.cliente.nombre_completo, f.estado, round(f.total, 2), round(costo_ag, 2), round(ganancia, 2)])

    ws4.column_dimensions['B'].width = 15
    ws4.column_dimensions['C'].width = 20
    ws4.column_dimensions['D'].width = 30
    for col in ['E', 'F', 'G', 'H']: ws4.column_dimensions[col].width = 18

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"Dashboard_{titulo_periodo}_{hoy.strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = make_response(buffer.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
'''

if 'exportar_excel' not in content:
    content += excel_endpoint

with open('routes/dashboard.py', 'w', encoding='utf-8') as f:
    f.write(content)

print("Dashboard routes updated successfully!")
