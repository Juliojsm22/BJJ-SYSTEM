from flask import Blueprint, render_template, jsonify, redirect, url_for, flash
from flask_login import login_required, current_user
from models import Cliente, Paquete, Factura, db, get_local_now
from datetime import datetime, timedelta
from sqlalchemy import func, case

dashboard_bp = Blueprint('dashboard', __name__, url_prefix='/dashboard')

from flask import request

from sqlalchemy.orm import joinedload

def calcular_ganancia_facturas(inicio, fin=None):
    from models import COSTOS_AGENCIA
    query = Factura.query.options(joinedload(Factura.paquetes)).filter(
        Factura.estado.in_(['finalizada', 'pagada']),
        Factura.fecha_emision >= inicio
    )
    if fin:
        query = query.filter(Factura.fecha_emision < fin)
    
    facturas = query.all()
    total_ganancia = 0.0
    for f in facturas:
        costo_agencia = 0
        for p in f.paquetes:
            origen = p.origen if getattr(p, 'origen', None) else 'miami'
            tarifas_origen = COSTOS_AGENCIA.get(origen, COSTOS_AGENCIA['miami'])
            tarifa_aplicada = tarifas_origen.get(p.tipo_envio, tarifas_origen['aereo'])
            costo_agencia += (p.peso * tarifa_aplicada)
            
        total_ganancia += (f.total - costo_agencia)
    return round(total_ganancia, 2)

@dashboard_bp.route('/')
@login_required
def index():
    if current_user.rol != 'admin':
        return redirect(url_for('paquetes.index'))
        
    periodo = request.args.get('periodo', 'mes')
    mes_historico = request.args.get('mes_historico')
    hoy = get_local_now()
    inicio_semana = (hoy - timedelta(days=hoy.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    inicio_mes = hoy.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    
    fin_periodo = None
    
    if periodo == 'historico' and mes_historico:
        year, month = map(int, mes_historico.split('-'))
        inicio_periodo = hoy.replace(year=year, month=month, day=1, hour=0, minute=0, second=0, microsecond=0)
        if month == 12:
            fin_periodo = hoy.replace(year=year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
        else:
            fin_periodo = hoy.replace(year=year, month=month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
    elif periodo == 'dia':
        inicio_periodo = hoy.replace(hour=0, minute=0, second=0, microsecond=0)
    elif periodo == 'semana':
        inicio_periodo = inicio_semana
    else:  # mes
        inicio_periodo = inicio_mes

    filtros_fecha = [Factura.fecha_emision >= inicio_periodo]
    if fin_periodo:
        filtros_fecha.append(Factura.fecha_emision < fin_periodo)

    ganancias_semana = calcular_ganancia_facturas(inicio_semana)
    if periodo == 'historico':
        ganancias_mes = calcular_ganancia_facturas(inicio_periodo, fin_periodo)
    else:
        ganancias_mes = calcular_ganancia_facturas(inicio_mes)

    total_clientes = Cliente.query.filter_by(activo=True).count()
    total_paquetes = Paquete.query.count()
    paquetes_sin_facturar = Paquete.query.filter_by(factura_id=None).count()
    facturas_pendientes = Factura.query.filter_by(estado='borrador').count()

    top_clientes = db.session.query(
        Cliente.nombre_completo,
        func.sum(Paquete.peso).label('total_libras'),
        func.count(Paquete.id).label('total_paquetes')
    ).join(Factura, Cliente.id == Factura.cliente_id)\
     .join(Paquete, Factura.id == Paquete.factura_id)\
     .filter(Cliente.activo == True, Factura.estado.in_(['finalizada', 'pagada']), *filtros_fecha)\
     .group_by(Cliente.id)\
     .order_by(func.sum(Paquete.peso).desc())\
     .limit(5).all()

    tendencia_labels = []
    tendencia_valores = []

    if periodo == 'dia':
        # Últimos 7 días
        for i in range(6, -1, -1):
            dia = hoy - timedelta(days=i)
            inicio = dia.replace(hour=0, minute=0, second=0, microsecond=0)
            fin = inicio + timedelta(days=1)
            total = calcular_ganancia_facturas(inicio, fin)
            
            lbl = inicio.strftime('%d %b')
            if i == 0: lbl = "Hoy"
            tendencia_labels.append(lbl)
            tendencia_valores.append(float(total))

    elif periodo == 'semana':
        # Últimas 4 semanas
        for i in range(3, -1, -1):
            inicio = inicio_semana - timedelta(days=7 * i)
            fin = inicio + timedelta(days=7)
            total = calcular_ganancia_facturas(inicio, fin)
            
            lbl = f"{inicio.strftime('%d %b')}"
            if i == 0: lbl = "Esta Sem"
            tendencia_labels.append(lbl)
            tendencia_valores.append(float(total))

    else:
        # Últimos 6 meses hasta inicio_periodo
        base_date = inicio_periodo if periodo == 'historico' else hoy
        for i in range(5, -1, -1):
            target_month = base_date.month - i
            target_year = base_date.year
            while target_month <= 0:
                target_month += 12
                target_year -= 1
                
            inicio = base_date.replace(year=target_year, month=target_month, day=1, hour=0, minute=0, second=0, microsecond=0)
            if target_month == 12:
                fin = base_date.replace(year=target_year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
            else:
                fin = base_date.replace(year=target_year, month=target_month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
                
            total = calcular_ganancia_facturas(inicio, fin)
            tendencia_labels.append(inicio.strftime('%b %Y'))
            tendencia_valores.append(float(total))

    aereos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'aereo',
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).count()

    maritimos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'maritimo',
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).count()

    # Legacy variables for tables
    meses_data = []
    semanas_data = []
    for i in range(len(tendencia_labels)):
        item = {'total': tendencia_valores[i]}
        if periodo == 'mes':
            item['mes'] = tendencia_labels[i]
            meses_data.append(item)
        elif periodo == 'semana':
            item['semana'] = tendencia_labels[i]
            semanas_data.append(item)
        else:
            item['dia'] = tendencia_labels[i]
            # Usar semanas_data para no romper la tabla de abajo, adaptándola en el HTML
            item['semana'] = tendencia_labels[i] 
            semanas_data.append(item)

    return render_template('dashboard/index.html',
        periodo=periodo,
        mes_historico=mes_historico,
        ganancias_semana=ganancias_semana,
        ganancias_mes=ganancias_mes,
        total_clientes=total_clientes,
        total_paquetes=total_paquetes,
        paquetes_sin_facturar=paquetes_sin_facturar,
        facturas_pendientes=facturas_pendientes,
        top_clientes=top_clientes,
        meses_data=meses_data,
        semanas_data=semanas_data,
        tendencia_labels=tendencia_labels,
        tendencia_valores=tendencia_valores,
        aereos=aereos,
        maritimos=maritimos
    )

@dashboard_bp.route('/api/stats')
@login_required
def api_stats():
    hoy = get_local_now()
    inicio_mes = hoy.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    ganancias_mes = calcular_ganancia_facturas(inicio_mes)
    return jsonify({'ganancias_mes': float(ganancias_mes)})

@dashboard_bp.route('/pdf-reporte')
@login_required
def pdf_reporte():
    if current_user.rol != 'admin':
        flash('No tienes permisos.', 'error')
        return redirect(url_for('paquetes.index'))
        
    from flask import make_response
    import io
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    import os

    periodo = request.args.get('periodo', 'mes')
    mes_historico = request.args.get('mes_historico')
    hoy = get_local_now()
    inicio_semana = (hoy - timedelta(days=hoy.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    
    if periodo == 'historico' and mes_historico:
        year, month = map(int, mes_historico.split('-'))
        inicio_periodo = hoy.replace(year=year, month=month, day=1, hour=0, minute=0, second=0, microsecond=0)
    else:
        inicio_periodo = hoy

    tendencia_data = []
    titulo_tabla = ""
    encabezado_columna = ""

    if periodo == 'dia':
        titulo_tabla = 'Resumen Diario (Últimos 7 Días)'
        encabezado_columna = 'Día'
        for i in range(6, -1, -1):
            dia = hoy - timedelta(days=i)
            inicio = dia.replace(hour=0, minute=0, second=0, microsecond=0)
            fin = inicio + timedelta(days=1)
            total = calcular_ganancia_facturas(inicio, fin)
            
            lbl = inicio.strftime('%d %b')
            if i == 0: lbl = "Hoy"
            tendencia_data.append([lbl, f'${total:.2f}'])

    elif periodo == 'semana':
        titulo_tabla = 'Resumen Semanal (Últimas 4 Semanas)'
        encabezado_columna = 'Semana'
        for i in range(3, -1, -1):
            inicio = inicio_semana - timedelta(days=7 * i)
            fin = inicio + timedelta(days=7)
            total = calcular_ganancia_facturas(inicio, fin)
            
            lbl = f"{inicio.strftime('%d %b')} - {(fin - timedelta(days=1)).strftime('%d %b')}"
            if i == 0: lbl = "Esta Semana"
            tendencia_data.append([lbl, f'${total:.2f}'])

    else:
        titulo_tabla = f'Resumen Histórico ({mes_historico})' if periodo == 'historico' else 'Resumen Mensual (Últimos 6 Meses)'
        encabezado_columna = 'Mes'
        base_date = inicio_periodo if periodo == 'historico' else hoy
        for i in range(5, -1, -1):
            target_month = base_date.month - i
            target_year = base_date.year
            while target_month <= 0:
                target_month += 12
                target_year -= 1
                
            inicio = base_date.replace(year=target_year, month=target_month, day=1, hour=0, minute=0, second=0, microsecond=0)
            if target_month == 12:
                fin = base_date.replace(year=target_year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
            else:
                fin = base_date.replace(year=target_year, month=target_month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
                
            total = calcular_ganancia_facturas(inicio, fin)
            tendencia_data.append([inicio.strftime('%b %Y'), f'${total:.2f}'])

    # Reverse data so most recent is at top, like in HTML
    tendencia_data.reverse()

    buffer = io.BytesIO()
    numero_reporte = f"REP-{hoy.strftime('%Y%m%d%H%M')}"
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=inch, bottomMargin=inch, title=f"Reporte Financiero {numero_reporte}", author="BJJ SYSTEM")
    styles = getSampleStyleSheet()
    story = []

    logo_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'static', 'logoazul.PNG')
    if os.path.exists(logo_path):
        header_left = Image(logo_path, width=3.0*inch, height=1.0*inch, kind='proportional')
        header_left.hAlign = 'CENTER'
        story.append(header_left)
    
    story.append(Spacer(1, 0.3*inch))
    story.append(Paragraph(f'<font size=18 color="#3d5ba0"><b>REPORTE FINANCIERO ({periodo.upper()})</b></font>', styles['Title']))
    story.append(Paragraph(f'<font size=10 color="#666666">Generado el: {hoy.strftime("%d/%m/%Y %H:%M")}</font>', styles['Title']))
    story.append(Spacer(1, 0.5*inch))

    story.append(Paragraph(f'<b>{titulo_tabla}</b>', styles['Heading3']))
    story.append(Spacer(1, 0.1*inch))
    t_data = Table([[encabezado_columna, 'Ganancia Neta']] + tendencia_data, colWidths=[4*inch, 2*inch])
    t_data.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#3d5ba0')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (1,0), (1,-1), 'RIGHT'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#dddddd')),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_data)

    doc.build(story)
    
    response = make_response(buffer.getvalue())
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = f'inline; filename=reporte-financiero-{hoy.strftime("%Y%m%d")}.pdf'
    return response

@dashboard_bp.route('/fix-timezones')
@login_required
def fix_timezones():
    if current_user.rol != 'admin':
        flash('No tienes permisos.', 'error')
        return redirect(url_for('dashboard.index'))
        
    from models import Pago, HistorialRastreo, RegistroActividad
    
    # Bandera para saber si ya se corrigió
    ya_corregido = Factura.query.filter(Factura.fecha_emision < datetime(2025, 1, 1)).first()
    if ya_corregido:
        flash('Las fechas ya fueron corregidas anteriormente.', 'info')
        return redirect(url_for('dashboard.index'))
    
    # Restar 6 horas a todos los registros
    for c in Cliente.query.all():
        if c.creado_en: c.creado_en -= timedelta(hours=6)
    
    for p in Paquete.query.all():
        if p.registrado_en: p.registrado_en -= timedelta(hours=6)
        
    for f in Factura.query.all():
        if f.fecha_emision: f.fecha_emision -= timedelta(hours=6)
        
    for p in Pago.query.all():
        if p.fecha_pago: p.fecha_pago -= timedelta(hours=6)
        
    for h in HistorialRastreo.query.all():
        if h.creado_en: h.creado_en -= timedelta(hours=6)
        
    for r in RegistroActividad.query.all():
        if r.fecha: r.fecha -= timedelta(hours=6)
        
    # Crear un registro viejo falso para marcar que ya se corrigió y evitar doble resta
    falso = Factura(numero='FIX-TZ', cliente_id=1, fecha_emision=datetime(2020, 1, 1))
    db.session.add(falso)
    
    db.session.commit()
    flash('Se han corregido las horas de todos los registros antiguos exitosamente.', 'success')
    return redirect(url_for('facturas.index'))

@dashboard_bp.route('/exportar-excel')
@login_required
def exportar_excel():
    if current_user.rol != 'admin':
        flash('No tienes permisos.', 'error')
        return redirect(url_for('paquetes.index'))
        
    from flask import make_response
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    periodo = request.args.get('periodo', 'mes')
    mes_historico = request.args.get('mes_historico')
    hoy = get_local_now()
    inicio_semana = (hoy - timedelta(days=hoy.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    inicio_mes = hoy.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    
    fin_periodo = None

    if periodo == 'historico' and mes_historico:
        year, month = map(int, mes_historico.split('-'))
        inicio_periodo = hoy.replace(year=year, month=month, day=1, hour=0, minute=0, second=0, microsecond=0)
        if month == 12:
            fin_periodo = hoy.replace(year=year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
        else:
            fin_periodo = hoy.replace(year=year, month=month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
        titulo_periodo = f'Histórico ({mes_historico})'
    elif periodo == 'dia':
        inicio_periodo = hoy.replace(hour=0, minute=0, second=0, microsecond=0)
        titulo_periodo = 'Diario'
    elif periodo == 'semana':
        inicio_periodo = inicio_semana
        titulo_periodo = 'Semanal'
    else:  # mes
        inicio_periodo = inicio_mes
        titulo_periodo = 'Mensual'
        
    filtros_fecha = [Factura.fecha_emision >= inicio_periodo]
    if fin_periodo:
        filtros_fecha.append(Factura.fecha_emision < fin_periodo)

    # 1. KPIs Globales
    ganancia_periodo = calcular_ganancia_facturas(inicio_periodo, fin_periodo)
    total_clientes = Cliente.query.filter_by(activo=True).count()
    total_paquetes = Paquete.query.count()
    aereos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'aereo',
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).count()
    maritimos = Paquete.query.join(Factura).filter(
        Paquete.tipo_envio == 'maritimo',
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).count()

    wb = openpyxl.Workbook()
    
    # --- HOJA 1: RESUMEN ---
    ws1 = wb.active
    ws1.title = "Resumen"
    header_fill = PatternFill(start_color='3D5BA0', end_color='3D5BA0', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')

    ws1.append([f"RESUMEN {titulo_periodo.upper()}"])
    ws1['A1'].font = Font(bold=True, size=14, color='3D5BA0')
    ws1.append([])
    
    resumen_data = [
        ["Indicador", "Valor"],
        ["Ganancia del Periodo ($)", round(ganancia_periodo, 2)],
        ["Paquetes Aéreos (Facturados)", aereos],
        ["Paquetes Marítimos (Facturados)", maritimos],
        ["Total Clientes Activos", total_clientes],
        ["Total Paquetes en Sistema", total_paquetes]
    ]
    
    for row_idx, row_data in enumerate(resumen_data, 3):
        ws1.append(row_data)
        if row_idx == 3:
            ws1.cell(row=row_idx, column=1).fill = header_fill
            ws1.cell(row=row_idx, column=1).font = header_font
            ws1.cell(row=row_idx, column=2).fill = header_fill
            ws1.cell(row=row_idx, column=2).font = header_font

    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['B'].width = 15

    # --- HOJA 2: GRAFICA ---
    ws2 = wb.create_sheet(title="Evolución (Gráfica)")
    ws2.append(["Periodo", "Ganancia Neta ($)"])
    ws2['A1'].fill = header_fill; ws2['A1'].font = header_font
    ws2['B1'].fill = header_fill; ws2['B1'].font = header_font

    if periodo == 'dia':
        for i in range(6, -1, -1):
            dia = hoy - timedelta(days=i)
            inicio = dia.replace(hour=0, minute=0, second=0, microsecond=0)
            fin = inicio + timedelta(days=1)
            total = calcular_ganancia_facturas(inicio, fin)
            lbl = inicio.strftime('%d %b')
            ws2.append([lbl if i != 0 else "Hoy", total])
    elif periodo == 'semana':
        for i in range(3, -1, -1):
            inicio = inicio_semana - timedelta(days=7 * i)
            fin = inicio + timedelta(days=7)
            total = calcular_ganancia_facturas(inicio, fin)
            lbl = f"{inicio.strftime('%d %b')} - {(fin - timedelta(days=1)).strftime('%d %b')}"
            ws2.append([lbl if i != 0 else "Esta Semana", total])
    else:
        base_date = inicio_periodo if periodo == 'historico' else hoy
        for i in range(5, -1, -1):
            target_month = base_date.month - i
            target_year = base_date.year
            while target_month <= 0:
                target_month += 12; target_year -= 1
            inicio = base_date.replace(year=target_year, month=target_month, day=1, hour=0, minute=0, second=0, microsecond=0)
            if target_month == 12: fin = base_date.replace(year=target_year+1, month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
            else: fin = base_date.replace(year=target_year, month=target_month+1, day=1, hour=0, minute=0, second=0, microsecond=0)
            total = calcular_ganancia_facturas(inicio, fin)
            ws2.append([inicio.strftime('%b %Y'), total])

    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 20

    # --- HOJA 3: TOP CLIENTES ---
    ws3 = wb.create_sheet(title="Top Clientes")
    ws3.append(["Posición", "Cliente", "Total Libras", "Total Paquetes"])
    for col in ['A', 'B', 'C', 'D']:
        ws3[f'{col}1'].fill = header_fill; ws3[f'{col}1'].font = header_font

    top_clientes = db.session.query(
        Cliente.nombre_completo,
        func.sum(Paquete.peso).label('total_libras'),
        func.count(Paquete.id).label('total_paquetes')
    ).join(Factura, Cliente.id == Factura.cliente_id).join(Paquete, Factura.id == Paquete.factura_id).filter(
        Cliente.activo == True, Factura.estado.in_(['finalizada', 'pagada']), *filtros_fecha
    ).group_by(Cliente.id).order_by(func.sum(Paquete.peso).desc()).limit(10).all()

    for idx, (nombre, libras, paquetes) in enumerate(top_clientes, 1):
        ws3.append([idx, nombre, float(libras or 0), paquetes])

    ws3.column_dimensions['B'].width = 30
    ws3.column_dimensions['C'].width = 15
    ws3.column_dimensions['D'].width = 15

    # --- HOJA 4: DETALLE FACTURAS ---
    ws4 = wb.create_sheet(title="Detalle Facturas")
    headers4 = ["Factura ID", "Número", "Fecha Emisión", "Cliente", "Estado", "Total Libras", "Total ($)", "Costo Agencia ($)", "Ganancia Neta ($)"]
    ws4.append(headers4)
    for i, _ in enumerate(headers4, 1):
        ws4.cell(row=1, column=i).fill = header_fill
        ws4.cell(row=1, column=i).font = header_font

    facturas = Factura.query.options(joinedload(Factura.paquetes)).filter(
        Factura.estado.in_(['finalizada', 'pagada']),
        *filtros_fecha
    ).order_by(Factura.fecha_emision.desc()).all()

    from models import COSTOS_AGENCIA
    for f in facturas:
        costo_ag = 0
        total_libras = 0
        for p in f.paquetes:
            if getattr(p, 'peso', None):
                total_libras += p.peso
            if getattr(p, 'categoria', 'general') in ['celular', 'laptop']:
                origen_key = p.origen if getattr(p, 'origen', None) else 'miami'
                costo_unidad = COSTOS_AGENCIA.get(origen_key, COSTOS_AGENCIA['miami']).get(p.categoria, 0)
                costo_ag += (p.cantidad or 1) * costo_unidad
            else:
                origen = p.origen if getattr(p, 'origen', None) else 'miami'
                tarifas_origen = COSTOS_AGENCIA.get(origen, COSTOS_AGENCIA['miami'])
                tipo = 'aereo' if str(p.tipo_envio).strip().lower() in ['aereo', 'aéreo'] else 'maritimo'
                costo_ag += (p.peso * tarifas_origen.get(tipo, tarifas_origen['aereo']))
        ganancia = f.total - costo_ag
        ws4.append([f.id, f.numero, f.fecha_emision.strftime('%Y-%m-%d %H:%M') if f.fecha_emision else '',
                    f.cliente.nombre_completo, f.estado, round(total_libras, 2), round(f.total, 2), round(costo_ag, 2), round(ganancia, 2)])

    ws4.column_dimensions['B'].width = 15
    ws4.column_dimensions['C'].width = 20
    ws4.column_dimensions['D'].width = 30
    for col in ['E', 'F', 'G', 'H', 'I']: ws4.column_dimensions[col].width = 18

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"Dashboard_{titulo_periodo}_{hoy.strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = make_response(buffer.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
